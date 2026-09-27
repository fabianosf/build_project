"""Normalize chat attachments (text / images / PDF / DOCX) for LLM prompts."""

from __future__ import annotations

import base64
import binascii
from io import BytesIO
from typing import Any

MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_ATTACH_TEXT_CHARS = 80_000
MAX_IMAGES = 3
# Cap when enriching suggest/forge request text
SUGGEST_ATTACH_CHARS = 12_000
FORGE_ATTACH_CHARS = 20_000

TEXT_MIMES = {
    "text/plain",
    "text/markdown",
    "text/csv",
    "application/json",
    "application/csv",
}
IMAGE_MIMES = {"image/jpeg", "image/png", "image/webp", "image/jpg"}
PDF_MIMES = {"application/pdf"}
DOCX_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/docx",
}
DOC_LEGACY_MIMES = {
    "application/msword",
    "application/doc",
}


def _decode_b64(data_b64: str) -> bytes:
    raw = (data_b64 or "").strip()
    if "," in raw and raw.lower().startswith("data:"):
        raw = raw.split(",", 1)[1]
    try:
        data = base64.b64decode(raw, validate=False)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("Anexo com base64 inválido.") from exc
    if len(data) > MAX_FILE_BYTES:
        raise ValueError(
            f"Anexo excede {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )
    return data


def _extract_pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError as exc:
        raise ValueError(
            "Suporte a PDF indisponível (instale pypdf). "
            "Cole o texto ou envie .txt/.md."
        ) from exc

    try:
        reader = PdfReader(BytesIO(data))
        parts: list[str] = []
        for page in reader.pages[:40]:
            parts.append(page.extract_text() or "")
        text = "\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Não foi possível ler o PDF: {exc}") from exc
    if not text:
        raise ValueError("PDF sem texto extraível (pode ser só imagem).")
    return text


def _extract_docx_text(data: bytes) -> str:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise ValueError(
            "Suporte a DOCX indisponível (instale python-docx). "
            "Envie .txt/.md/.pdf."
        ) from exc

    try:
        doc = Document(BytesIO(data))
        parts: list[str] = []
        for para in doc.paragraphs:
            t = (para.text or "").strip()
            if t:
                parts.append(t)
        for table in doc.tables:
            for row in table.rows:
                cells = [
                    (c.text or "").strip() for c in row.cells if (c.text or "").strip()
                ]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts).strip()
    except Exception as exc:  # noqa: BLE001
        raise ValueError(f"Não foi possível ler o DOCX: {exc}") from exc
    if not text:
        raise ValueError("DOCX sem texto extraível.")
    return text


def process_attachments(
    attachments: list[dict[str, Any]] | None,
) -> tuple[str, list[dict[str, Any]], list[str]]:
    """Return (text_block, image_parts, warnings)."""
    if not attachments:
        return "", [], []

    text_chunks: list[str] = []
    image_parts: list[dict[str, Any]] = []
    warnings: list[str] = []
    text_budget = MAX_ATTACH_TEXT_CHARS

    for item in attachments:
        if not isinstance(item, dict):
            continue
        name = str(item.get("name") or "anexo")[:200]
        lower_name = name.lower()
        mime = str(item.get("mime") or "application/octet-stream").lower().strip()
        if mime == "image/jpg":
            mime = "image/jpeg"

        if lower_name.endswith(".doc") and not lower_name.endswith(".docx"):
            warnings.append(
                f"{name}: .doc legado não suportado — envie .docx ou PDF."
            )
            continue
        if mime in DOC_LEGACY_MIMES and not lower_name.endswith(".docx"):
            warnings.append(
                f"{name}: .doc legado não suportado — envie .docx ou PDF."
            )
            continue

        if mime in TEXT_MIMES or lower_name.endswith(
            (".txt", ".md", ".json", ".csv")
        ):
            text = item.get("text")
            if text is None and item.get("data_base64"):
                try:
                    text = _decode_b64(str(item["data_base64"])).decode(
                        "utf-8", errors="replace"
                    )
                except ValueError as exc:
                    warnings.append(f"{name}: {exc}")
                    continue
            text = str(text or "")
            if len(text) > text_budget:
                text = text[:text_budget]
                warnings.append(f"{name}: texto truncado ao limite.")
            text_budget -= len(text)
            text_chunks.append(f"### Anexo: {name}\n{text}")
            continue

        if mime in PDF_MIMES or lower_name.endswith(".pdf"):
            try:
                data = _decode_b64(str(item.get("data_base64") or ""))
                extracted = _extract_pdf_text(data)
            except ValueError as exc:
                warnings.append(f"{name}: {exc}")
                continue
            if len(extracted) > text_budget:
                extracted = extracted[:text_budget]
                warnings.append(f"{name}: PDF truncado ao limite.")
            text_budget -= len(extracted)
            text_chunks.append(f"### Anexo PDF: {name}\n{extracted}")
            continue

        if mime in DOCX_MIMES or lower_name.endswith(".docx"):
            try:
                data = _decode_b64(str(item.get("data_base64") or ""))
                extracted = _extract_docx_text(data)
            except ValueError as exc:
                warnings.append(f"{name}: {exc}")
                continue
            if len(extracted) > text_budget:
                extracted = extracted[:text_budget]
                warnings.append(f"{name}: DOCX truncado ao limite.")
            text_budget -= len(extracted)
            text_chunks.append(f"### Anexo DOCX: {name}\n{extracted}")
            continue

        if mime in IMAGE_MIMES or lower_name.endswith(
            (".jpg", ".jpeg", ".png", ".webp")
        ):
            if len(image_parts) >= MAX_IMAGES:
                warnings.append(f"{name}: limite de {MAX_IMAGES} imagens.")
                continue
            b64 = str(item.get("data_base64") or "").strip()
            if not b64:
                warnings.append(f"{name}: imagem sem dados.")
                continue
            try:
                _decode_b64(b64)  # size check
            except ValueError as exc:
                warnings.append(f"{name}: {exc}")
                continue
            if b64.startswith("data:"):
                url = b64
            else:
                url = f"data:{mime};base64,{b64}"
            image_parts.append(
                {"type": "image_url", "image_url": {"url": url}}
            )
            continue

        warnings.append(f"{name}: tipo não suportado ({mime}).")

    block = ""
    if text_chunks:
        block = "## Anexos do usuário\n\n" + "\n\n".join(text_chunks)
    return block, image_parts, warnings


def enrich_request_text(
    request_text: str,
    attachments: list[dict[str, Any]] | None,
    *,
    max_attach_chars: int = SUGGEST_ATTACH_CHARS,
) -> tuple[str, list[str], int]:
    """Merge attachment text into request for suggest/forge.

    Returns (enriched_text, warnings, image_count).
    """
    base = (request_text or "").strip()
    block, images, warnings = process_attachments(attachments)
    warn = list(warnings)
    if images:
        warn.append(
            f"{len(images)} imagem(ns) anexada(s) — "
            "usadas no chat se o modelo aceitar visão."
        )
    if block:
        clipped = block
        if len(clipped) > max_attach_chars:
            clipped = clipped[: max_attach_chars - 40] + "\n\n[...anexos truncados...]\n"
            warn.append("Texto dos anexos truncado para análise do pedido.")
        enriched = f"{base}\n\n{clipped}".strip() if base else clipped
    else:
        enriched = base
    return enriched, warn, len(images)


def build_user_content(
    message: str,
    attach_text: str,
    image_parts: list[dict[str, Any]],
) -> str | list[dict[str, Any]]:
    body = message.strip()
    if attach_text:
        body = f"{body}\n\n{attach_text}" if body else attach_text
    if not image_parts:
        return body
    parts: list[dict[str, Any]] = [
        {"type": "text", "text": body or "(imagem anexada)"}
    ]
    parts.extend(image_parts)
    return parts
